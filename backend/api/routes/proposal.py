"""
提案生成 API 路由
================

處理研究提案的生成、編輯和修訂功能
"""

from fastapi import APIRouter, HTTPException, BackgroundTasks
from fastapi.responses import FileResponse
from pydantic import BaseModel, validator
from typing import List, Optional, Dict, Any
import asyncio
import sys
import os
import tempfile
import io
import requests
from docx import Document as DocxDocument
from docx.shared import Inches
from io import BytesIO
import logging
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, HTTPException, UploadFile, File, Form

# 配置 logger
logger = logging.getLogger(__name__)
from fastapi.responses import FileResponse
import tempfile
import os
import re
from docx import Document as DocxDocument
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json

# 添加原項目路徑到 sys.path
sys.path.append(os.path.join(os.path.dirname(__file__), '../../../app'))

# 修正導入方式
try:
    from backend.services.pubchem_service import chemical_metadata_extractor
except ImportError:
    # 如果直接導入失敗，嘗試使用完整路徑 (已重組，不再需要)
    # sys.path.append(os.path.join(os.path.dirname(__file__), '../../../'))
    from backend.services.pubchem_service import chemical_metadata_extractor
from backend.services.chemical_service import chemical_service
from backend.core import demo_config
from backend.services import demo_service
from langchain_core.documents import Document

# SVG 轉換依賴檢查
try:
    from svglib.svglib import svg2rlg
    from reportlab.graphics import renderPDF
    SVGLIB_AVAILABLE = True
except ImportError as e:
    SVGLIB_AVAILABLE = False

# PyMuPDF 用於 PDF 到 PNG 轉換
try:
    import fitz
    PYMUPDF_AVAILABLE = True
except ImportError as e:
    PYMUPDF_AVAILABLE = False

import time
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from PIL import Image

router = APIRouter()

class ProposalRequest(BaseModel):
    """提案生成請求模型"""
    research_goal: str
    user_feedback: Optional[str] = None
    previous_proposal: Optional[str] = None
    retrieval_count: Optional[int] = 10  # 預設檢索 10 個文檔
    mof_linker_mode: Optional[str] = "auto"  # "auto", "single", "dual"

    @validator('research_goal')
    def validate_research_goal(cls, v):
        """驗證研究目標"""
        if not v or not v.strip():
            raise ValueError('研究目標不能為空')
        if len(v.strip()) < 3:
            raise ValueError('研究目標至少需要3個字符')
        if len(v.strip()) > 10000:
            raise ValueError('研究目標不能超過10000個字符')
        return v.strip()

    @validator('retrieval_count')
    def validate_retrieval_count(cls, v):
        """驗證檢索數量"""
        if v is not None:
            if v < 0:
                raise ValueError('檢索數量不能小於0')
            if v > 100:
                raise ValueError('檢索數量不能超過100')
        return v

class ProposalResponse(BaseModel):
    """提案生成響應模型"""
    proposal: str
    chemicals: List[Dict[str, Any]]
    experiment_detail: Optional[str] = None
    citations: List[Dict[str, str]]
    not_found: List[str]
    # 以可序列化的結構回傳 chunks：[{ page_content, metadata }]
    chunks: List[Dict[str, Any]]
    used_model: Optional[str] = None
    structured_proposal: Optional[Dict[str, Any]] = None


class ProposalRevisionRequest(BaseModel):
    """提案修訂請求模型"""
    original_proposal: str
    user_feedback: str
    # 來自前端的可序列化 chunks
    chunks: List[Dict[str, Any]]
    # 可選的檢索參數
    k_new_chunks: Optional[int] = 3  # 新chunks檢索數量，預設3（降低查詢量）
    mof_linker_mode: Optional[str] = "auto"  # "auto", "single", "dual"


def _serialize_chunks(chunks: List[Any]) -> List[Dict[str, Any]]:
    """將 LangChain Document 物件序列化為可回傳的 dict 結構。"""
    serialized: List[Dict[str, Any]] = []
    for doc in chunks or []:
        try:
            serialized.append({
                "page_content": getattr(doc, "page_content", ""),
                "metadata": getattr(doc, "metadata", {}) or {}
            })
        except Exception:
            continue
    return serialized


def _deserialize_chunks(chunks_like: List[Dict[str, Any]]) -> List[Document]:
    """將前端傳來的 dict 結構還原為 LangChain Document。"""
    documents: List[Document] = []
    for item in chunks_like or []:
        page_content = item.get("page_content", "")
        metadata = item.get("metadata", {}) or {}
        documents.append(Document(page_content=page_content, metadata=metadata))
    return documents


def _resolve_proposal_names_to_smiles(structured_proposal: Dict[str, Any] | None, chemical_metadata_list: List[Dict[str, Any]]) -> None:
    """
    將 structured_proposal 中的 mof_linker_name 與 mof_linker_name_2
    解析為 SMILES 並寫入 mof_linker_smiles 與 mof_linker_smiles_2，
    以便前端可以直接讀取使用。
    """
    if not structured_proposal:
        return

    l1_name = structured_proposal.get("mof_linker_name", "")
    l2_name = structured_proposal.get("mof_linker_name_2", "")

    def metadata_smiles(linker_name: str) -> str:
        target = linker_name.strip().casefold()
        if not target:
            return ""
        for chemical in chemical_metadata_list:
            aliases = {
                chemical.get("name"),
                chemical.get("iupac_name"),
                chemical.get("query_name"),
            }
            if target in {
                str(alias).strip().casefold()
                for alias in aliases
                if alias
            }:
                return chemical.get("smiles") or ""
        return ""

    def resolve_smiles(linker_name: str) -> str:
        if not linker_name:
            return ""
        matched = metadata_smiles(linker_name)
        if matched:
            return matched
        try:
            from backend.services.mof.pormake_resolver import resolve_linker_input

            smiles, _provenance = resolve_linker_input(linker_name)
            return smiles
        except Exception as exc:
            print(f"⚠️ 查詢 linker SMILES 失敗 ({linker_name}): {exc}")
            return ""

    l1_smiles = resolve_smiles(l1_name)
    l2_smiles = resolve_smiles(l2_name)

    # 3. 寫入 JSON 以便保持對前端接口的兼容性
    structured_proposal["mof_linker_smiles"] = l1_smiles
    structured_proposal["mof_linker_smiles_2"] = l2_smiles

@router.post("/proposal/generate", response_model=ProposalResponse)
async def generate_proposal(request: ProposalRequest):
    """
    生成研究提案

    Args:
        request: 包含研究目標的請求

    Returns:
        生成的提案內容，包括化學品信息和實驗細節
    """
    if demo_config.is_stage_demo("proposal"):
        return ProposalResponse(**demo_service.get_proposal_response(request.research_goal))

    import time
    import uuid

    # 生成唯一的請求 ID
    request_id = str(uuid.uuid4())[:8]
    start_time = time.time()

    print(f"🚀 [DEBUG-{request_id}] ========== 開始處理提案生成請求 ==========")
    print(f"🚀 [DEBUG-{request_id}] 時間戳: {time.strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"🚀 [DEBUG-{request_id}] 收到請求 research_goal = '{request.research_goal}'")
    print(f"🚀 [DEBUG-{request_id}] retrieval_count = {request.retrieval_count}")
    print(f"🚀 [DEBUG-{request_id}] 請求來源: {request}")

    try:
        print(f"🔍 [DEBUG-{request_id}] 準備調用 agent_answer with mode='make proposal'")

        # 延遲導入以避免循環導入問題
        from backend.services.knowledge_service import agent_answer

        # 與 Streamlit Tab1 對齊：使用模式 make proposal 生成提案
        with demo_config.stage_context("proposal"):
            result = agent_answer(
                request.research_goal,
                mode="make proposal",
                k=request.retrieval_count,
                mof_linker_mode=request.mof_linker_mode,
            )

        print(f"🔍 [DEBUG-{request_id}] agent_answer 調用成功")
        print(f"🔍 [DEBUG-{request_id}] result 類型: {type(result)}")
        print(f"🔍 [DEBUG-{request_id}] result 鍵: {list(result.keys())}")
        print(f"🔍 [DEBUG-{request_id}] result['answer'] 長度: {len(result.get('answer', ''))}")
        print(f"🔍 [DEBUG-{request_id}] result['answer'] 內容: {result.get('answer', '')[:200]}...")

        from backend.services.pubchem_service import remove_json_chemical_block

        structured_proposal = result.get("structured_proposal")
        proposal_answer = remove_json_chemical_block(result.get("answer", ""))

        # 第二階段 MOF extraction 不依賴 Chemical Summary。先完成這條路徑，
        # 避免不乾淨的 materials_list 或 PubChem 暫時失敗阻塞 PORMAKE。
        from backend.core.generation import extract_mof_from_proposal

        print(f"🔍 [DEBUG-{request_id}] 啟動第二階段 AI，從提案文本中提取 MOF 金屬與配體")
        # Extraction can make a second LLM call after agent_answer; preserve the
        # same request stage so an unrelated enabled Demo checkbox cannot block it.
        with demo_config.stage_context("proposal"):
            extraction_res = extract_mof_from_proposal(
                proposal_answer,
                mof_linker_mode=request.mof_linker_mode,
                chemicals_list=None,
            )
        if structured_proposal is None:
            structured_proposal = {}
        if extraction_res.get("is_mof_related", False):
            structured_proposal["mof_metal_element"] = extraction_res.get(
                "mof_metal_element", ""
            )
            structured_proposal["mof_linker_name"] = extraction_res.get(
                "mof_linker_name", ""
            )
            structured_proposal["mof_linker_name_2"] = extraction_res.get(
                "mof_linker_name_2", ""
            )
        else:
            structured_proposal["mof_metal_element"] = ""
            structured_proposal["mof_linker_name"] = ""
            structured_proposal["mof_linker_name_2"] = ""

        _resolve_proposal_names_to_smiles(structured_proposal, [])

        # Chemical Summary is best-effort and cannot invalidate MOF extraction.
        print(f"🔍 [DEBUG-{request_id}] 準備調用化學服務提取化學品並添加結構圖")
        chemical_metadata_list = []
        not_found_list = []
        try:
            if structured_proposal.get("materials_list"):
                print(f"🔍 [DEBUG-{request_id}] 使用結構化數據中的材料列表: {structured_proposal['materials_list']}")
                from backend.services.pubchem_service import extract_and_fetch_chemicals

                chemical_metadata_list, not_found_list = extract_and_fetch_chemicals(
                    structured_proposal["materials_list"]
                )

                print(f"🔍 [DEBUG-{request_id}] 為結構化數據的化學品添加SMILES繪製")
                print(f"🔍 [DEBUG-{request_id}] 化學品數量: {len(chemical_metadata_list)}")

                enhanced_chemicals = []
                for i, chemical in enumerate(chemical_metadata_list):
                    print(f"🔍 [DEBUG-{request_id}] 處理化學品 {i+1}/{len(chemical_metadata_list)}: {chemical.get('name', 'Unknown')}")
                    enhanced_chemical = chemical_service.add_smiles_drawing(chemical)
                    enhanced_chemicals.append(enhanced_chemical)
                chemical_metadata_list = enhanced_chemicals
            else:
                print(f"🔍 [DEBUG-{request_id}] 回退到從文本中提取材料列表")
                (
                    chemical_metadata_list,
                    not_found_list,
                    _cleaned_answer,
                ) = chemical_service.extract_chemicals_with_drawings(
                    result.get("answer", "")
                )
        except Exception as exc:
            print(
                f"⚠️ [DEBUG-{request_id}] Chemical Summary 失敗，"
                f"保留 MOF extraction 結果: {exc}"
            )
            chemical_metadata_list = []
            not_found_list = []

        print(f"🔍 [DEBUG-{request_id}] 化學品提取和結構圖生成完成")
        print(f"🔍 [DEBUG-{request_id}] proposal_answer 長度: {len(proposal_answer)}")
        print(f"🔍 [DEBUG-{request_id}] chemical_metadata_list 數量: {len(chemical_metadata_list)}")

        # Chemical Summary 可能提供更精確的 query_name/IUPAC alias，再解析一次。
        _resolve_proposal_names_to_smiles(structured_proposal, chemical_metadata_list)

        citations = result.get("citations", [])
        chunks = result.get("chunks", [])
        used_model = result.get("used_model", "unknown")

        # 修復 citations 中的 page 欄位類型問題
        fixed_citations = []
        for citation in citations:
            fixed_citation = citation.copy()
            # 確保 page 欄位是字串
            if "page" in fixed_citation:
                fixed_citation["page"] = str(fixed_citation["page"])
            fixed_citations.append(fixed_citation)

        end_time = time.time()
        duration = end_time - start_time

        print(f"✅ [DEBUG-{request_id}] ========== 提案生成完成 ==========")
        print(f"✅ [DEBUG-{request_id}] 總耗時: {duration:.2f} 秒")
        print(f"✅ [DEBUG-{request_id}] 檢索到的文檔數量: {len(chunks)}")
        print(f"✅ [DEBUG-{request_id}] 引用數量: {len(fixed_citations)}")
        print(f"✅ [DEBUG-{request_id}] 化學品數量: {len(chemical_metadata_list)}")

        return ProposalResponse(
            proposal=proposal_answer,
            chemicals=chemical_metadata_list,
            citations=fixed_citations,
            not_found=not_found_list,
            chunks=_serialize_chunks(chunks),
            used_model=used_model,
            structured_proposal=structured_proposal
        )

    except Exception as e:
        end_time = time.time()
        duration = end_time - start_time
        print(f"❌ [DEBUG-{request_id}] ========== 提案生成失敗 ==========")
        print(f"❌ [DEBUG-{request_id}] 總耗時: {duration:.2f} 秒")
        print(f"❌ [DEBUG-{request_id}] 錯誤: {str(e)}")
        import traceback
        traceback.print_exc()

        # 根據錯誤類型返回適當的HTTP狀態碼
        error_message = str(e)
        if "Connection error" in error_message or "API" in error_message:
            # API連接錯誤，返回503服務不可用
            raise HTTPException(
                status_code=503,
                detail="AI服務暫時不可用，請稍後再試"
            )
        elif "validation" in error_message.lower() or "invalid" in error_message.lower():
            # 驗證錯誤，返回400錯誤請求
            raise HTTPException(
                status_code=400,
                detail=f"請求參數錯誤: {error_message}"
            )
        else:
            # 其他錯誤，返回500內部服務器錯誤
            raise HTTPException(
                status_code=500,
                detail=f"提案生成失敗: {error_message}"
            )

@router.post("/proposal/revise", response_model=ProposalResponse)
async def revise_proposal(request: ProposalRevisionRequest):
    """
    根據用戶反饋修訂提案

    Args:
        request: 包含原始提案和用戶反饋的請求

    Returns:
        修訂後的提案內容
    """
    if demo_config.is_stage_demo("generate_new_idea"):
        return ProposalResponse(**demo_service.get_revision_response(request.user_feedback))

    try:
        # 延遲導入以避免循環導入問題
        from backend.services.knowledge_service import agent_answer

        # 檢查開發模式狀態
        from backend.core.settings_manager import settings_manager
        is_dev_mode = settings_manager.get_dev_mode_status()

        # 根據開發模式決定檢索參數
        k_new_chunks = 1 if is_dev_mode else (request.k_new_chunks or 3)

        # 與 Streamlit Tab1 對齊：採用 generate new idea 模式，並帶入原始提案與 chunks
        with demo_config.stage_context("generate_new_idea"):
            result = agent_answer(
                request.user_feedback,
                mode="generate new idea",
                old_chunks=_deserialize_chunks(request.chunks),
                proposal=request.original_proposal,
                k_new_chunks=k_new_chunks,  # 傳遞檢索參數
                mof_linker_mode=request.mof_linker_mode,
            )

        from backend.services.pubchem_service import remove_json_chemical_block

        proposal_answer = remove_json_chemical_block(result.get("answer", ""))
        structured_proposal = result.get("structured_proposal")
        if structured_proposal is None:
            structured_proposal = {}

        # 先完成 MOF extraction 與 linker resolution；Chemical Summary 為
        # best-effort，不得阻塞 PORMAKE 路徑。
        from backend.core.generation import extract_mof_from_proposal

        print(f"🔍 [DEBUG] 啟動第二階段 AI，從修訂提案文本中提取 MOF 金屬與配體")
        # Extraction can make a second LLM call after agent_answer; preserve the
        # request stage so an unrelated enabled Demo checkbox cannot block it.
        with demo_config.stage_context("generate_new_idea"):
            extraction_res = extract_mof_from_proposal(
                proposal_answer,
                mof_linker_mode=request.mof_linker_mode,
                chemicals_list=None,
            )
        if extraction_res.get("is_mof_related", False):
            structured_proposal["mof_metal_element"] = extraction_res.get("mof_metal_element", "")
            structured_proposal["mof_linker_name"] = extraction_res.get("mof_linker_name", "")
            structured_proposal["mof_linker_name_2"] = extraction_res.get("mof_linker_name_2", "")
        else:
            structured_proposal["mof_metal_element"] = ""
            structured_proposal["mof_linker_name"] = ""
            structured_proposal["mof_linker_name_2"] = ""

        _resolve_proposal_names_to_smiles(structured_proposal, [])

        material_names = (
            result.get("materials_list")
            or structured_proposal.get("materials_list")
            or []
        )
        chemical_metadata_list = []
        not_found_list = []
        try:
            if material_names:
                print(f"🔍 [DEBUG] 使用結構化數據中的材料列表: {material_names}")
                from backend.services.pubchem_service import extract_and_fetch_chemicals

                chemical_metadata_list, not_found_list = extract_and_fetch_chemicals(
                    material_names
                )
            else:
                print(f"🔍 [DEBUG] 回退到從文本中提取材料列表")
                (
                    chemical_metadata_list,
                    not_found_list,
                    _cleaned_answer,
                ) = chemical_metadata_extractor(result.get("answer", ""))

            print(f"🔍 [DEBUG] 為修訂提案的化學品添加SMILES繪製")
            chemical_metadata_list = [
                chemical_service.add_smiles_drawing(chemical)
                for chemical in chemical_metadata_list
            ]
        except Exception as exc:
            print(f"⚠️ [DEBUG] Chemical Summary 失敗，保留 MOF extraction: {exc}")
            chemical_metadata_list = []
            not_found_list = []

        _resolve_proposal_names_to_smiles(structured_proposal, chemical_metadata_list)

        # 修復 citations 中的 page 欄位類型問題
        fixed_citations = []
        for citation in result.get("citations", []):
            fixed_citation = citation.copy()
            # 確保 page 欄位是字串
            if "page" in fixed_citation:
                fixed_citation["page"] = str(fixed_citation["page"])
            fixed_citations.append(fixed_citation)

        return ProposalResponse(
            proposal=proposal_answer,
            chemicals=chemical_metadata_list,
            citations=fixed_citations,
            not_found=not_found_list,
            chunks=_serialize_chunks(result.get("chunks", [])),
            structured_proposal=structured_proposal
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"提案修訂失敗: {str(e)}")

class ExperimentDetailRequest(BaseModel):
    """生成實驗細節請求模型"""
    proposal: str
    chunks: List[Dict[str, Any]]


@router.post("/proposal/experiment-detail")
async def generate_experiment_detail(request: ExperimentDetailRequest):
    """
    生成實驗細節

    Args:
        proposal: 提案內容
        chunks: 相關文檔片段

    Returns:
        實驗細節內容
    """
    if demo_config.is_stage_demo("experiment_detail"):
        return demo_service.get_experiment_detail_response(request.proposal)

    try:
        # 延遲導入以避免循環導入問題
        from backend.services.knowledge_service import agent_answer

        # 與 Streamlit Tab1 對齊：由 agent 以指定模式展開實驗細節
        with demo_config.stage_context("experiment_detail"):
            result = agent_answer(
                "",
                mode="expand to experiment detail",
                chunks=_deserialize_chunks(request.chunks),
                proposal=request.proposal,
            )

        # ✅ 修復：添加citations字段到返回結果
        # 從result中獲取citations，如果沒有則返回空列表
        citations = result.get("citations", [])
        print(f"🔍 [DEBUG] generate_experiment_detail 返回的citations數量: {len(citations)}")

        return {
            "experiment_detail": result.get("answer", ""),
            "structured_experiment": result.get("structured_experiment", {}),
            "citations": citations,  # ✅ 修復：添加citations字段
            "success": True,
            "retry_info": {
                "retry_count": getattr(result, 'retry_count', 0),
                "final_tokens": getattr(result, 'final_tokens', 0)
            }
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"實驗細節生成失敗: {str(e)}")

@router.get("/proposal/status/{task_id}")
async def get_proposal_status(task_id: str):
    """
    獲取提案生成任務狀態

    Args:
        task_id: 任務 ID

    Returns:
        任務狀態信息
    """
    # TODO: 實現任務狀態追蹤
    return {
        "task_id": task_id,
        "status": "completed",
        "progress": 100
    }

def clean_text_for_xml(text):
    """清理文本以確保XML兼容性"""
    if not text:
        return ""
    # 移除NULL字節和控制字符
    cleaned = "".join(char for char in text if ord(char) >= 32 or char in '\n\r\t')
    # 移除其他可能導致XML問題的字符
    cleaned = cleaned.replace('\x00', '')  # NULL字節
    cleaned = cleaned.replace('\x01', '')  # SOH
    cleaned = cleaned.replace('\x02', '')  # STX
    cleaned = cleaned.replace('\x03', '')  # ETX
    cleaned = cleaned.replace('\x04', '')  # EOT
    cleaned = cleaned.replace('\x05', '')  # ENQ
    cleaned = cleaned.replace('\x06', '')  # ACK
    cleaned = cleaned.replace('\x07', '')  # BEL
    cleaned = cleaned.replace('\x08', '')  # BS
    cleaned = cleaned.replace('\x0b', '')  # VT
    cleaned = cleaned.replace('\x0c', '')  # FF
    cleaned = cleaned.replace('\x0e', '')  # SO
    cleaned = cleaned.replace('\x0f', '')  # SI
    return cleaned

def clean_markdown_text(text):
    """清理 markdown 格式，轉換為純文本"""
    if not text:
        return ""

    import re

    # 移除 markdown 格式
    cleaned = text
    cleaned = re.sub(r'\*\*(.*?)\*\*', r'\1', cleaned)  # 移除粗體標記
    cleaned = re.sub(r'\*(.*?)\*', r'\1', cleaned)  # 移除斜體標記
    cleaned = re.sub(r'`(.*?)`', r'\1', cleaned)  # 移除代碼標記
    cleaned = re.sub(r'^#+\s*(.*)$', r'\1', cleaned, flags=re.MULTILINE)  # 移除標題標記
    cleaned = re.sub(r'^\s*[-*+]\s+', '- ', cleaned, flags=re.MULTILINE)  # 統一項目符號
    cleaned = re.sub(r'^\s*\d+\.\s+', '', cleaned, flags=re.MULTILINE)  # 移除編號
    cleaned = re.sub(r'\n\s*\n\s*\n', '\n\n', cleaned)  # 移除多餘空行
    cleaned = re.sub(r'\n\s*\*\*', '\n', cleaned)  # 移除粗體前的換行
    cleaned = re.sub(r'\*\*\s*\n', '\n', cleaned)  # 移除粗體後的換行

    return cleaned

def get_nfpa_icon_image_stream(nfpa_code: str) -> io.BytesIO:
    """
    用 Headless Chrome 抓取 PubChem NFPA 圖，回傳 BytesIO PNG 圖片
    """
    nfpa_svg_url = f"https://pubchem.ncbi.nlm.nih.gov/image/nfpa.cgi?code={nfpa_code}"

    chrome_options = Options()
    chrome_options.add_argument("--headless=new")
    chrome_options.add_argument("--disable-gpu")
    chrome_options.add_argument("--window-size=300,300")
    chrome_options.add_argument("--hide-scrollbars")

    driver = webdriver.Chrome(options=chrome_options)
    driver.get(nfpa_svg_url)
    time.sleep(1)

    png_data = driver.get_screenshot_as_png()
    driver.quit()

    image = Image.open(io.BytesIO(png_data))
    cropped = image.crop((0, 0, 100, 100))  # 你已經測試好的範圍
    output = io.BytesIO()
    cropped.save(output, format="PNG")
    output.seek(0)
    return output

class DocxRequest(BaseModel):
    """DOCX 生成請求模型"""
    proposal: str
    chemicals: List[Dict[str, Any]]
    not_found: List[str]
    experiment_detail: Optional[str] = ""
    citations: List[Dict[str, Any]]  # ✅ 修復：改為 Any 以支持數字類型的 page 字段

@router.post("/proposal/generate-docx")
async def generate_docx(request: DocxRequest):
    """
    生成 DOCX 文件

    Args:
        request: 包含所有提案數據的請求

    Returns:
        DOCX 文件下載響應
    """
    tmp_path = None
    try:
        print(f"🔍 BACKEND DEBUG: 開始生成 DOCX 文件")
        print(f"🔍 BACKEND DEBUG: proposal 長度: {len(request.proposal)}")
        print(f"🔍 BACKEND DEBUG: chemicals 數量: {len(request.chemicals)}")
        print(f"🔍 BACKEND DEBUG: experiment_detail 長度: {len(request.experiment_detail)}")
        print(f"🔍 BACKEND DEBUG: citations 數量: {len(request.citations)}")

        doc = DocxDocument()
        doc.add_heading("AI Generated Research Proposal", 0)

        # Proposal Section
        doc.add_heading("Proposal", level=1)
        proposal_text = clean_text_for_xml(clean_markdown_text(request.proposal))
        doc.add_paragraph(proposal_text)

        # Chemical Table
        doc.add_heading("Chemical Summary Table", level=1)
        table = doc.add_table(rows=1, cols=8)  # 多兩欄：Structure + Safety
        hdr = table.rows[0].cells
        hdr[0].text = "Structure"
        hdr[1].text = "Name"
        hdr[2].text = "Formula"
        hdr[3].text = "MW"
        hdr[4].text = "Boiling Point (°C)"
        hdr[5].text = "Melting Point (°C)"
        hdr[6].text = "CAS No."
        hdr[7].text = "Safety Icons"

        for chem in request.chemicals:
            row = table.add_row().cells

            # 優先使用 SMILES 繪製的結構圖
            if chem.get("png_structure"):
                try:
                    # 從 Base64 轉換為圖片流
                    import base64
                    img_data = base64.b64decode(chem["png_structure"].split(",")[1])
                    img_stream = BytesIO(img_data)
                    row[0].paragraphs[0].add_run().add_picture(img_stream, width=Inches(1))
                    print(f"✅ 使用 SMILES 繪製的結構圖: {chem.get('name', 'Unknown')}")
                except Exception as e:
                    print(f"⚠️ SMILES 圖片插入失敗: {chem.get('name', 'Unknown')}, {e}")
                    row[0].text = "SMILES image error"
            elif chem.get("image_url"):
                # 備用方案：使用原有的 URL 圖片
                try:
                    response = requests.get(chem["image_url"], timeout=5)
                    if response.status_code == 200:
                        img_stream = BytesIO(response.content)
                        row[0].paragraphs[0].add_run().add_picture(img_stream, width=Inches(1))
                        print(f"✅ 使用 URL 圖片: {chem.get('name', 'Unknown')}")
                    else:
                        row[0].text = "Image not found"
                        print(f"⚠️ URL 圖片下載失敗: {chem.get('name', 'Unknown')}")
                except Exception as e:
                    print(f"⚠️ 圖片下載失敗: {chem['image_url']}, {e}")
                    row[0].text = "Image error"
            else:
                row[0].text = "No image"
                print(f"⚠️ 沒有圖片資料: {chem.get('name', 'Unknown')}")

            # 文字欄位 - 使用清理函數
            row[1].text = clean_text_for_xml(chem.get("name", "-") or "-")
            row[2].text = clean_text_for_xml(chem.get("formula", "-") or "-")
            row[3].text = clean_text_for_xml(str(chem.get("weight", "-") or "-"))
            row[4].text = clean_text_for_xml(str(chem.get("boiling_point_c", "-") or "-"))
            row[5].text = clean_text_for_xml(str(chem.get("melting_point_c", "-") or "-"))
            row[6].text = clean_text_for_xml(chem.get("cas", "-") or "-")

            # Safety icons（支援多張 GHS + 1 張 NFPA）
            icons_cell = row[7].paragraphs[0]
            ghs_icons = chem.get("safety_icons", {}).get("ghs_icons", [])
            nfpa_icon_url = chem.get("safety_icons", {}).get("nfpa_image")

            for icon_url in ghs_icons:
                try:
                    icon_resp = requests.get(icon_url, timeout=5)
                    if icon_resp.status_code == 200:
                        # 使用 svglib + PyMuPDF 轉換 SVG 為 PNG
                        if SVGLIB_AVAILABLE and PYMUPDF_AVAILABLE:
                            try:
                                # 步驟1：SVG 轉 PDF
                                drawing = svg2rlg(io.BytesIO(icon_resp.content))
                                if drawing:
                                    # 創建臨時 PDF 文件
                                    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_pdf:
                                        renderPDF.drawToFile(drawing, tmp_pdf.name)

                                    # 步驟2：PDF 轉 PNG
                                    pdf_document = fitz.open(tmp_pdf.name)
                                    page = pdf_document[0]
                                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))  # 2x 縮放

                                    # 將 PNG 轉換為 BytesIO
                                    png_stream = io.BytesIO()
                                    png_stream.write(pix.tobytes("png"))
                                    png_stream.seek(0)

                                    # 插入到 Word 文檔
                                    run = icons_cell.add_run()
                                    run.add_picture(png_stream, width=Inches(0.3))

                                    # 清理臨時文件
                                    pdf_document.close()
                                    os.unlink(tmp_pdf.name)

                                else:
                                    print(f"⚠️ SVG 轉換失敗 - drawing 為 None: {icon_url}")
                            except Exception as e:
                                print(f"⚠️ SVG 轉換錯誤: {icon_url}, {e}")
                        else:
                            print(f"⚠️ svglib 或 PyMuPDF 不可用，無法轉換 SVG: {icon_url}")
                except Exception as e:
                    print(f"⚠️ Failed to convert or insert icon: {icon_url}, {e}")

            if nfpa_icon_url:
                try:
                    # 從 URL 擷取出 NFPA code，例如 https://...code=130
                    from urllib.parse import urlparse, parse_qs
                    parsed = urlparse(nfpa_icon_url)
                    nfpa_code = parse_qs(parsed.query).get("code", [""])[0]

                    if nfpa_code:
                        image_stream = get_nfpa_icon_image_stream(nfpa_code)
                        if image_stream:
                            run = icons_cell.add_run()
                            run.add_picture(image_stream, width=Inches(0.3))
                    else:
                        print(f"⚠️ 無法從 URL 擷取 NFPA code: {nfpa_icon_url}")
                except Exception as e:
                    print(f"⚠️ Failed to convert or insert NFPA icon: {nfpa_icon_url}, {e}")

        # Not Found Chemicals
        if request.not_found:
            doc.add_heading("Not Found Chemicals", level=2)
            for name in request.not_found:
                doc.add_paragraph(f"- {clean_text_for_xml(name)}")

        # Experiment Details
        doc.add_heading("Experimental Plan", level=1)
        experiment_text = clean_text_for_xml(clean_markdown_text(request.experiment_detail))
        doc.add_paragraph(experiment_text)

        # Citations
        doc.add_heading("Citations", level=1)
        for i, c in enumerate(request.citations, 1):
            title = clean_text_for_xml(c.get('title', ''))
            page = clean_text_for_xml(str(c.get('page', '')))
            snippet = clean_text_for_xml(c.get('snippet', ''))
            doc.add_paragraph(f"[{i}] {title} | Page {page} | Snippet: {snippet}")

        # Save and Download
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            doc.save(tmp.name)
            tmp_path = tmp.name
            logger.info(f"DOCX 文件已保存到: {tmp_path}")

        logger.info("準備返回 FileResponse")
        return FileResponse(
            path=tmp_path,
            filename="proposal_report.docx",
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    except Exception as e:
        logger.error(f"DOCX 生成失敗: {str(e)}", exc_info=True)
        # 清理臨時文件
        if 'tmp_path' in locals() and tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except OSError as cleanup_error:
                logger.warning(f"清理臨時文件失敗: {cleanup_error}")
        raise HTTPException(status_code=500, detail=f"DOCX 生成失敗: {str(e)}")

@router.post("/proposal/test-docx")
async def test_docx_generation():
    """
    測試 DOCX 生成功能
    """
    try:
        logger.info("開始測試 DOCX 生成")

        doc = DocxDocument()
        doc.add_heading("Test Document", 0)
        doc.add_paragraph("This is a test document to verify DOCX generation works.")

        # Save and Download
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            doc.save(tmp.name)
            tmp_path = tmp.name
            logger.info(f"測試 DOCX 文件已保存到: {tmp_path}")

        return FileResponse(
            path=tmp_path,
            filename="test_document.docx",
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    except Exception as e:
        logger.error(f"測試 DOCX 生成失敗: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"測試 DOCX 生成失敗: {str(e)}")
